import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- DATA LOGBOOK HARIAN ---
DAILY_LOGS = [
    {
        "no": 1,
        "hari_tanggal": "Senin, 10 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 1: Analisis Kebutuhan",
        "modul": "Orientasi & Problem Scoping",
        "kegiatan": "Orientasi di lingkungan Yayasan Gereja Protestan Kampung Bali (YGPKB). Wawancara awal dengan Kepala Departemen (KD) Sarpras terkait kendala pelaporan manual fasilitas fisik, tiket keluhan yang kerap tercecer di grup percakapan WhatsApp, dan ketiadaan riwayat terpusat untuk perbaikan sarpras di 14 unit yayasan.",
        "output": "Catatan observasi lapangan, rumusan masalah sistem (Problem Statement), dan matriks pemetaan unit operasional YGPKB.",
        "status": "Selesai"
    },
    {
        "no": 2,
        "hari_tanggal": "Selasa, 11 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 1: Analisis Kebutuhan",
        "modul": "Wawancara Lapangan ASP Wilayah",
        "kegiatan": "Wawancara dan observasi kerja bersama Area Supervisor (ASP) Wilayah Kubu Raya dan Wilayah Reguler. Menganalisis alur verifikasi fisik fasilitas ke lokasi sekolah, penaksiran estimasi biaya darurat, serta penanganan jadwal servis preventif rutin (AC, pompa air, genset) di kompleks sekolah.",
        "output": "Diagram alur kerja manual survei lapangan oleh tim ASP dan identifikasi kebutuhan modul pemeliharaan preventif berkala.",
        "status": "Selesai"
    },
    {
        "no": 3,
        "hari_tanggal": "Rabu, 12 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 1: Analisis Kebutuhan",
        "modul": "Analisis Prosedur PO & BAST Sarpras",
        "kegiatan": "Wawancara dengan staf Admin Sarpras dan Bagian Purchasing. Mengumpulkan berkas fisik dan formulir manual pengajuan unit, Purchase Order (PO), serta format Berita Acara Serah Terima (BAST) untuk 5 jenis pekerjaan (Perbaikan, Pemeliharaan, Pengadaan, Mutasi, Pemutihan).",
        "output": "Kumpulan arsip format dokumen fisik resmi YGPKB dan spesifikasi kebutuhan atribut data pada dokumen PO dan BAST.",
        "status": "Selesai"
    },
    {
        "no": 4,
        "hari_tanggal": "Kamis, 13 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 1: Analisis Kebutuhan",
        "modul": "Spesifikasi Hak Akses & SRS",
        "kegiatan": "Mengidentifikasi dan memetakan hak akses sistem ke dalam 7 Point of View (POV): Kepala Sekolah/Pimpinan ITBSS, ASP Wilayah, KD Sarpras, Admin Sarpras, Koordinator Utama/SKI, Dept Kantor Pusat, dan Super Admin. Menyusun draf Software Requirements Specification (SRS).",
        "output": "Dokumen draf SRS v1.0 dan Matriks Hak Akses Pengguna (Role-Based Access Control / RBAC).",
        "status": "Selesai"
    },
    {
        "no": 5,
        "hari_tanggal": "Jumat, 14 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 1: Analisis Kebutuhan",
        "modul": "Validasi & Finalisasi Scope",
        "kegiatan": "Rapat evaluasi bersama KD Sarpras dan Pembimbing Lapangan. Mempresentasikan draf alur bisnis sistem usulan SIPARAS, mendiskusikan batasan persetujuan anggaran berjenjang, dan memvalidasi ruang lingkup (scope) sistem berbasis arsitektur monorepo web application.",
        "output": "Notula rapat evaluasi kebutuhan dan dokumen spesifikasi kebutuhan sistem yang telah disetujui (Approved SRS).",
        "status": "Selesai"
    },
    {
        "no": 6,
        "hari_tanggal": "Selasa, 18 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 2: Perancangan Sistem",
        "modul": "Pemodelan UML & Alur Proses Bisnis",
        "kegiatan": "Membuat diagram Use Case untuk 7 peran pengguna, Activity Diagram untuk siklus hidup tiket dari formulir pengajuan hingga terbit BAST, serta Sequence Diagram untuk interaksi antara Vue frontend client dan Laravel RESTful API saat proses pengajuan dan approval berjenjang.",
        "output": "Dokumen rancangan UML lengkap (Use Case, Activity Diagram, dan Sequence Diagram alur tiket).",
        "status": "Selesai"
    },
    {
        "no": 7,
        "hari_tanggal": "Rabu, 19 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 2: Perancangan Sistem",
        "modul": "Pemodelan Basis Data (ERD)",
        "kegiatan": "Merancang struktur basis data relasional. Menentukan entitas utama: users, units, rooms, vendors, tickets, ticket_items, ticket_photos, purchase_orders, po_items, bast_documents, preventive_schedules, dan audit_logs. Melakukan normalisasi hingga bentuk normal ketiga (3NF).",
        "output": "Diagram Entity-Relationship (ERD) konseptual dan relasi database ter-normalisasi (3NF).",
        "status": "Selesai"
    },
    {
        "no": 8,
        "hari_tanggal": "Kamis, 20 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 2: Perancangan Sistem",
        "modul": "Perancangan Skema DDL & Kamus Data",
        "kegiatan": "Menyusun spesifikasi Data Definition Language (DDL) spesifik untuk MySQL 8.4 LTS InnoDB. Menentukan presisi tipe data, foreign keys cascading, indeks komposit untuk optimasi query pencarian tiket, serta struktur kamus data (Data Dictionary).",
        "output": "Dokumen kamus data terperinci dan berkas skrip DDL SQL database SIPARAS.",
        "status": "Selesai"
    },
    {
        "no": 9,
        "hari_tanggal": "Jumat, 21 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 2: Perancangan Sistem",
        "modul": "Wireframe & UI/UX Design Mockup",
        "kegiatan": "Merancang wireframe antarmuka pengguna (UI/UX) untuk desktop dan browser smartphone: halaman login bernuansa korporat, dashboard modular, formulir tiket dinamis (step wizard), linimasa visual (timeline tracker), dan pratinjau lembar cetak dokumen PO/BAST.",
        "output": "Wireframe mockup desain antarmuka pengguna SIPARAS (Desktop & Mobile view).",
        "status": "Selesai"
    },
    {
        "no": 10,
        "hari_tanggal": "Senin, 24 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 3: Backend Core",
        "modul": "Setup Repositori & Environment Laravel 11",
        "kegiatan": "Menginisialisasi repositori Git dengan struktur monorepo. Melakukan setup backend framework Laravel 11 (PHP 8.3/8.4), konfigurasi environment .env, instalasi dependency via Composer, dan konfigurasi koneksi database engine MySQL 8.4 LTS lokal.",
        "output": "Proyek backend Laravel 11 terkonfigurasi pada repositori lokal dan koneksi basis data aktif.",
        "status": "Selesai"
    },
    {
        "no": 11,
        "hari_tanggal": "Selasa, 25 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 3: Backend Core",
        "modul": "Database Migrations",
        "kegiatan": "Menulis skrip database migration Laravel untuk 13 tabel utama: tabel master data (units, rooms, vendors), tabel transaksional (tickets, ticket_items, ticket_photos), tabel pengadaan (purchase_orders, po_items), tabel BAST, kalender preventif, dan audit_logs.",
        "output": "13 file migration Laravel berhasil dieksekusi tanpa kendala (php artisan migrate:fresh).",
        "status": "Selesai"
    },
    {
        "no": 12,
        "hari_tanggal": "Rabu, 26 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 3: Backend Core",
        "modul": "Database Seeders & Account Fixtures",
        "kegiatan": "Mengembangkan database seeder (DatabaseSeeder.php, MasterSeeder.php, UserSeeder.php). Menginput data awal 14 unit kerja YGPKB, master ruangan baku, serta akun uji coba untuk 13 Kepala Sekolah, 5 ASP Wilayah, KD Sarpras, Admin Sarpras, dan Koordinator ke dalam dokumentasi account.txt.",
        "output": "Database terisi data awal testing dan dokumentasi daftar kredensial akun uji coba (account.txt).",
        "status": "Selesai"
    },
    {
        "no": 13,
        "hari_tanggal": "Kamis, 27 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 3: Backend Core",
        "modul": "Autentikasi API (Laravel Sanctum)",
        "kegiatan": "Mengonfigurasi library Laravel Sanctum untuk autentikasi berbasis Bearer Token. Membuat AuthController.php dengan endpoint login, logout, dan me (profil pengguna aktif beserta relasi peran dan unit kerja). Melakukan pengujian request API via Postman.",
        "output": "Endpoint API /api/v1/auth/login, /logout, /me berjalan dengan validasi token yang aman.",
        "status": "Selesai"
    },
    {
        "no": 14,
        "hari_tanggal": "Jumat, 28 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 3: Backend Core",
        "modul": "Role-Based Access Control (RBAC) Middleware",
        "kegiatan": "Mengembangkan middleware kustom RoleMiddleware.php untuk memvalidasi hak akses endpoint API berdasarkan 7 peran pengguna. Mengimplementasikan isolasi data: Kepala Sekolah hanya dapat mengakses unitnya, ASP hanya mengakses unit dalam wilayah tugasnya, serta akses global bagi Admin dan KD Sarpras.",
        "output": "Middleware RBAC aktif dan teruji membatasi hak akses data secara ketat antar peran.",
        "status": "Selesai"
    },
    {
        "no": 15,
        "hari_tanggal": "Senin, 31 Agt 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 4: Backend Logic & Service",
        "modul": "API Master Data & Ruangan Cascading",
        "kegiatan": "Mengembangkan MasterDataController.php untuk pengelolaan data unit, ruangan (Room), dan rekanan/vendor. Menambahkan logika cascading dropdown (mengambil daftar ruangan berdasarkan unit sekolah terpilih) serta endpoint penambahan ruangan cepat (quick-add room).",
        "output": "RESTful API Master Data lengkap dengan filter relasional cascading unit-ruangan.",
        "status": "Selesai"
    },
    {
        "no": 16,
        "hari_tanggal": "Selasa, 1 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 4: Backend Logic & Service",
        "modul": "Modul Pengajuan Tiket (Ticket API)",
        "kegiatan": "Mengembangkan TicketController.php dan form request validation StoreTicketRequest.php. Mengimplementasikan penyimpanan multi-item (TicketItem), penomoran tiket unik terotomatisasi sesuai kode unit dan tanggal, serta validasi atribut untuk 5 jenis prosedur sarpras.",
        "output": "Endpoint POST /api/v1/tickets dan fungsi pembuatan tiket multi-item selesai diimplementasikan.",
        "status": "Selesai"
    },
    {
        "no": 17,
        "hari_tanggal": "Rabu, 2 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 4: Backend Logic & Service",
        "modul": "Service Layer Approval & Audit Logs",
        "kegiatan": "Mengembangkan ApprovalRoutingService.php dan AuditLogService.php. Mengatur logika transisi status tiket (Draft -> Submitted -> ASP Review -> KD Sarpras Approval -> In Progress -> Completed/Rejected) serta pencatatan otomatis setiap riwayat aksi, nama pengguna, dan timestamp ke tabel audit_logs.",
        "output": "Logika alur persetujuan berjenjang otomatis dan perekaman jejak audit status tiket terpasang.",
        "status": "Selesai"
    },
    {
        "no": 18,
        "hari_tanggal": "Kamis, 3 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 4: Backend Logic & Service",
        "modul": "Manajemen Purchase Order & Terbilang Rupiah",
        "kegiatan": "Mengembangkan PurchaseOrderController.php, PurchaseOrderService.php, dan helper konversi angka ke nominal kalimat teks Indonesia (TerbilangHelper.php). Mengintegrasikan pembuatan PO otomatis ketika tiket sarpras telah mendapatkan persetujuan anggaran dari KD Sarpras/Koordinator.",
        "output": "Endpoint manajemen PO dan fungsi konversi otomatis nominal rupiah ke kalimat resmi bahasa Indonesia.",
        "status": "Selesai"
    },
    {
        "no": 19,
        "hari_tanggal": "Jumat, 4 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 4: Backend Logic & Service",
        "modul": "Engine Dokumen PDF (DomPDF) & Export Excel",
        "kegiatan": "Mengintegrasikan library barryvdh/laravel-dompdf untuk pembuatan berkas PDF resmi PO dan 5 varian BAST berkop resmi YGPKB melalui BastDocumentController.php. Mengintegrasikan maatwebsite/excel dengan class TicketsReportExport.php untuk ekspor rekapitulasi tiket sarpras ke format Excel.",
        "output": "Fitur generator dokumen cetak PDF (PO & BAST) serta fitur ekspor laporan ke format Excel (.xlsx).",
        "status": "Selesai"
    },
    {
        "no": 20,
        "hari_tanggal": "Senin, 7 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 5: Frontend UI/UX",
        "modul": "Setup Vue 3 + Tailwind CSS + Monorepo",
        "kegiatan": "Menginisialisasi frontend client menggunakan Vue 3 (Composition API) dan build tool Vite. Mengonfigurasi Tailwind CSS untuk utilitas antarmuka modern, tipografi Outfit/Inter, serta mengatur skrip npm run dev pada package.json root monorepo menggunakan concurrently untuk dual-server otomatis.",
        "output": "Frontend Vue 3 + Vite berjalan pada port 5173 terhubung secara monorepo dengan backend Laravel.",
        "status": "Selesai"
    },
    {
        "no": 21,
        "hari_tanggal": "Selasa, 8 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 5: Frontend UI/UX",
        "modul": "Routing, State Management & Axios Interceptor",
        "kegiatan": "Mengonfigurasi Vue Router (router/index.js) dengan navigation guards untuk proteksi rute halaman privat. Mengonfigurasi Pinia Store (stores/auth.js, ticket.js, master.js) dan membuat instance HTTP client Axios (services/api.js) dengan interceptor penyisipan Bearer Token dan penanganan 401 unauthenticated.",
        "output": "Sistem routing aman dengan manajemen state terpusat dan penanganan sesi autentikasi otomatis.",
        "status": "Selesai"
    },
    {
        "no": 22,
        "hari_tanggal": "Rabu, 9 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 5: Frontend UI/UX",
        "modul": "Tata Letak Dashboard & Navigasi Berbasis Role",
        "kegiatan": "Membangun layout antarmuka utama (DashboardLayout.vue), bilah navigasi atas (Navbar), serta bilah menu samping (AppSidebar.vue). Menyusun logika filter item menu samping secara dinamis agar hanya memunculkan menu yang relevan dengan hak akses (role) pengguna yang sedang aktif.",
        "output": "Layout dashboard terpadu dengan menu navigasi dinamis berbasis hak akses peran.",
        "status": "Selesai"
    },
    {
        "no": 23,
        "hari_tanggal": "Kamis, 10 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 5: Frontend UI/UX",
        "modul": "Halaman Login & Responsivitas Mobile",
        "kegiatan": "Merancang dan memprogram antarmuka halaman login (LoginView.vue). Mengoptimalkan tampilan pada perangkat layar kecil dengan proporsi viewport height (1/3 visual branding gradien YGPKB dan 2/3 formulir input kredensial) guna mencegah overflow dan pemotongan tampilan pada peramban smartphone.",
        "output": "Halaman login modern, responsif, dan bebas kendala tata letak pada perangkat mobile.",
        "status": "Selesai"
    },
    {
        "no": 24,
        "hari_tanggal": "Jumat, 11 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 5: Frontend UI/UX",
        "modul": "Antarmuka Pembuatan Tiket (Form Wizard)",
        "kegiatan": "Membangun halaman pembuatan tiket sarpras (TicketCreateView.vue). Membuat sub-komponen SmartCascadingLocation.vue untuk pemilihan unit dan ruangan serta ItemRepeaterTable.vue untuk memfasilitasi penginputan banyak daftar inventaris sekaligus dalam satu nomor tiket.",
        "output": "Formulir pengajuan tiket terstruktur dengan input repeater dinamis dan validasi formulir.",
        "status": "Selesai"
    },
    {
        "no": 25,
        "hari_tanggal": "Senin, 14 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 6: Komponen & Integrasi",
        "modul": "Kompresi Gambar & Integrasi Kamera Lapangan",
        "kegiatan": "Membangun komponen ImageCompressorUploader.vue. Mengimplementasikan kompresi gambar berbasis HTML5 Canvas pada sisi browser sebelum proses upload guna menghemat kuota dan mempercepat transmisi data. Menambahkan atribut capture='environment' untuk mengaktifkan akses kamera ponsel langsung bagi petugas lapangan.",
        "output": "Komponen upload bukti foto fisik dengan kompresi otomatis dan akses langsung kamera smartphone.",
        "status": "Selesai"
    },
    {
        "no": 26,
        "hari_tanggal": "Selasa, 15 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 6: Komponen & Integrasi",
        "modul": "Antarmuka Verifikasi ASP & Meja Approval KD",
        "kegiatan": "Membangun antarmuka verifikasi untuk ASP Wilayah (AspReviewQueueView.vue) dan antarmuka persetujuan untuk Kepala Departemen Sarpras (KdApprovalDeskView.vue). Menyediakan tombol aksi cepat: setujui (approve), tolak (reject), dan minta revisi formulir beserta dialog catatan perbaikan.",
        "output": "Antarmuka review dan persetujuan berjenjang selesai diintegrasikan dengan API approval.",
        "status": "Selesai"
    },
    {
        "no": 27,
        "hari_tanggal": "Rabu, 16 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 6: Komponen & Integrasi",
        "modul": "Rincian Tiket & Visual Timeline Tracker",
        "kegiatan": "Memprogram halaman rincian tiket (TicketDetailView.vue). Mengembangkan komponen visual linimasa (TimelineTracker.vue) untuk menampilkan tahapan perjalanan pengajuan secara transparan (Submitted -> Review ASP -> Approved KD -> In Progress -> BAST Selesai) lengkap dengan tanggal, jam, dan penanggung jawab aksi.",
        "output": "Tampilan detail tiket lengkap dengan tracker linimasa visual interaktif.",
        "status": "Selesai"
    },
    {
        "no": 28,
        "hari_tanggal": "Kamis, 17 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 6: Komponen & Integrasi",
        "modul": "Modul PO, Master Data & Laporan Eksekutif",
        "kegiatan": "Membangun antarmuka daftar Purchase Order (PoListView.vue), manajemen Master Data barang dan rekanan vendor (MasterDataView.vue, MasterVendorView.vue), manajemen akun pengguna (UserManagementView.vue), serta dasbor ringkasan status tiket sarpras (ExecutiveReportView.vue).",
        "output": "Antarmuka operasional inventaris, rekanan vendor, manajemen akun, dan ringkasan eksekutif.",
        "status": "Selesai"
    },
    {
        "no": 29,
        "hari_tanggal": "Jumat, 18 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 6: Komponen & Integrasi",
        "modul": "Komponen Modal, Toast & Navigasi Drawer Mobile",
        "kegiatan": "Mengembangkan komponen dialog konfirmasi global (ConfirmModal.vue, useConfirm.js), notifikasi pop-up (ToastNotification.vue, useToast.js), dan menyempurnakan navigasi drawer geser samping dengan tombol hamburger menu untuk kemudahan penggunaan pada ponsel petugas di lapangan.",
        "output": "Pengalaman pengguna (UX) responsif, dialog konfirmasi interaktif, dan navigasi mobile drawer yang mulus.",
        "status": "Selesai"
    },
    {
        "no": 30,
        "hari_tanggal": "Senin, 21 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 7: Testing & Deployment",
        "modul": "Pengujian Fungsional (Blackbox Testing)",
        "kegiatan": "Melakukan pengujian sistem secara komprehensif menggunakan metode Blackbox Testing pada seluruh modul: pengujian login 7 peran, isolasi data unit (Kepala Sekolah tidak bisa melihat tiket unit lain), validasi formulir input, alur approval berjenjang, dan kebenaran kalkulasi nominal pada modul PO.",
        "output": "Lembar matriks hasil pengujian Blackbox Testing (seluruh 20 skenario uji berhasil/Valid).",
        "status": "Selesai"
    },
    {
        "no": 31,
        "hari_tanggal": "Selasa, 22 Sep 2026",
        "jam": "08.00 - 16.30",
        "tahap": "Tahap 7: Testing & Deployment",
        "modul": "Deployment Cloud & Finalisasi Laporan",
        "kegiatan": "Melakukan konfigurasi build production frontend dan deployment ke platform hosting Vercel (https://siparas-frontend.vercel.app). Menghubungkan frontend ke API backend, melakukan verifikasi akhir cetak dokumen BAST/PO PDF berkop YGPKB, serta menyusun panduan penggunaan sistem (User Manual) dan laporan akhir pengerjaan.",
        "output": "Aplikasi SIPARAS live online dan siap dioperasikan, disertai dokumen laporan akhir dan logbook lengkap.",
        "status": "Selesai"
    }
]

# --- DATA REKAPITULASI MINGGUAN ---
WEEKLY_LOGS = [
    {
        "minggu": "Minggu 1",
        "tanggal": "10 - 14 Agt 2026",
        "fokus": "Analisis Kebutuhan & Observasi Lapangan",
        "capaian": "Penyelesaian wawancara stakeholder (KD Sarpras, ASP Wilayah, Kepala Sekolah, Admin Sarpras), pemetaan proses bisnis 5 prosedur sarpras, identifikasi 7 role pengguna, dan pengesahan dokumen SRS.",
        "kendala": "Proses birokrasi manual yayasan sangat beragam antar unit persekolahan (TK hingga SMA) dan kampus ITBSS, sehingga terdapat perbedaan format formulir fisik.",
        "solusi": "Melakukan standardisasi formulir dan mengelompokkan jenis pengajuan menjadi 5 prosedur baku yayasan (Perbaikan, Pemeliharaan, Pengadaan, Mutasi, Pemutihan).",
        "paraf": "Disetujui"
    },
    {
        "minggu": "Minggu 2",
        "tanggal": "18 - 21 Agt 2026",
        "fokus": "Perancangan Sistem & Arsitektur Database",
        "capaian": "Pembuatan diagram UML (Use Case, Activity, Sequence), perancangan Entity-Relationship Diagram (ERD) bentuk 3NF, penyusunan skema DDL MySQL 8.4 LTS, dan perancangan wireframe antarmuka UI/UX.",
        "kendala": "Kebutuhan relasi data yang kompleks karena satu tiket pengajuan sarpras dapat mencakup banyak barang inventaris, foto bukti sebelum/sesudah pengerjaan, dan riwayat penandatangan berbeda.",
        "solusi": "Mendesain tabel perantara relasional (ticket_items, ticket_photos, audit_logs) dengan foreign keys terindeks untuk menjaga integritas data.",
        "paraf": "Disetujui"
    },
    {
        "minggu": "Minggu 3",
        "tanggal": "24 - 28 Agt 2026",
        "fokus": "Setup Lingkungan & Backend Core",
        "capaian": "Inisialisasi arsitektur monorepo, instalasi Laravel 11, penulisan 13 berkas migration database, pembuatan seeder master data 14 unit dan akun uji coba, implementasi autentikasi Sanctum, dan middleware RBAC.",
        "kendala": "Risiko kebocoran data antar unit jika otorisasi hanya memfilter level pengguna umum tanpa memperhitungkan relasi wilayah kerja ASP dan keterikatan unit Kepala Sekolah.",
        "solusi": "Membangun middleware kustom RoleMiddleware yang secara ketat membatasi scope kueri Eloquent berdasarkan unit_id atau wilayah kerja ASP dari user token.",
        "paraf": "Disetujui"
    },
    {
        "minggu": "Minggu 4",
        "tanggal": "31 Agt - 4 Sep 2026",
        "fokus": "Logika Bisnis Backend & Service Layer",
        "capaian": "Pengembangan API Master Data (Unit, Room cascading, Vendor), API Pengajuan Tiket, ApprovalRoutingService, AuditLogService, integrasi modul Purchase Order dengan Terbilang Rupiah, serta generator PDF DomPDF & Excel.",
        "kendala": "Formatting cetak dokumen PDF resmi (BAST & PO) sering terpotong antar halaman saat rincian daftar barang melebihi batas tinggi kertas A4 standar.",
        "solusi": "Mengoptimalkan aturan CSS @page dan table page-break-inside: avoid pada view Blade Laravel DomPDF serta menyusun header tabel dinamis berulang.",
        "paraf": "Disetujui"
    },
    {
        "minggu": "Minggu 5",
        "tanggal": "7 - 11 Sep 2026",
        "fokus": "Pembangunan Frontend Client (Vue 3 & Pinia)",
        "capaian": "Inisialisasi Vue 3 Vite, styling Tailwind CSS, manajemen state Pinia, integrasi Axios interceptor, pembuatan layout dashboard adaptif, halaman login responsif, dan form pembuatan tiket dinamis.",
        "kendala": "Tampilan login pada layar ponsel mengalami pergeseran tata letak (overflow) saat keyboard virtual smartphone muncul.",
        "solusi": "Menerapkan pembagian proporsi visual tinggi layar (viewport height) dinamis: 1/3 area branding gradien YGPKB dan 2/3 formulir input dengan scrolling yang fleksibel.",
        "paraf": "Disetujui"
    },
    {
        "minggu": "Minggu 6",
        "tanggal": "14 - 18 Sep 2026",
        "fokus": "Komponen Interaktif, Optimasi Mobile & Integrasi UI",
        "capaian": "Pembuatan komponen ImageCompressorUploader dengan akses kamera langsung, antarmuka review ASP dan approval KD Sarpras, komponen visual linimasa TimelineTracker, modul PO, Master Data, dan modal konfirmasi.",
        "kendala": "Ukuran foto bukti fisik dari kamera smartphone beresolusi tinggi (3-8 MB) menyebabkan proses unggah data tiket di lapangan menjadi sangat lambat.",
        "solusi": "Mengembangkan kompresi gambar berbasis HTML5 Canvas pada sisi browser klien untuk mereduksi resolusi dan ukuran file hingga di bawah 400 KB sebelum dikirim ke API.",
        "paraf": "Disetujui"
    },
    {
        "minggu": "Minggu 7",
        "tanggal": "21 - 22 Sep 2026",
        "fokus": "Pengujian Sistem, Blackbox Testing & Deployment",
        "capaian": "Eksekusi pengujian fungsional Blackbox Testing pada 7 peran pengguna (seluruh 20 skenario berstatus valid), konfigurasi build produksi, deployment ke cloud Vercel, dan penyusunan buku panduan sistem.",
        "kendala": "Penanganan perbedaan domain (CORS) antara antarmuka frontend di platform Vercel dan backend server lokal/staging.",
        "solusi": "Mengatur konfigurasi cors.php di Laravel backend dan menyetel Base URL Axios pada environment variable Vite (VITE_API_BASE_URL).",
        "paraf": "Disetujui"
    }
]

# --- DATA BLACKBOX TESTING (20 SKENARIO) ---
TEST_CASES = [
    {"id": "TC-01", "modul": "Autentikasi", "skenario": "Login dengan kredensial valid (Kepala Sekolah)", "hasil_diharapkan": "Login berhasil, sistem menyimpan token Sanctum, redirect ke dashboard unit sekolah terkait.", "status": "Valid"},
    {"id": "TC-02", "modul": "Autentikasi", "skenario": "Login dengan password salah atau username tidak terdaftar", "hasil_diharapkan": "Sistem menolak autentikasi dan memunculkan notifikasi error kredensial tidak cocok.", "status": "Valid"},
    {"id": "TC-03", "modul": "RBAC & Otorisasi", "skenario": "Kepala Sekolah mencoba mengakses menu User Management", "hasil_diharapkan": "Sistem memblokir akses dan menampilkan halaman Forbidden (403 Access Denied).", "status": "Valid"},
    {"id": "TC-04", "modul": "RBAC & Otorisasi", "skenario": "Kepala Sekolah TK 1 mencoba melihat data tiket pengajuan SD 1", "hasil_diharapkan": "Sistem mengisolasi data, tiket unit lain tidak muncul pada tabel maupun via URL langsung.", "status": "Valid"},
    {"id": "TC-05", "modul": "RBAC & Otorisasi", "skenario": "ASP Wilayah Kota mengakses tiket di luar wilayahnya (misal Kubu Raya)", "hasil_diharapkan": "Sistem menyaring daftar antrean tiket hanya menampilkan unit-unit di Wilayah Kota.", "status": "Valid"},
    {"id": "TC-06", "modul": "Pembuatan Tiket", "skenario": "Kepala Sekolah mengisi form pengajuan perbaikan fasilitas", "hasil_diharapkan": "Nomor tiket ter-generate otomatis, item tersimpan, status tiket beralih ke 'Submitted'.", "status": "Valid"},
    {"id": "TC-07", "modul": "Pembuatan Tiket", "skenario": "Submit form tiket tanpa melampirkan data barang dan foto kendala", "hasil_diharapkan": "Validasi form aktif, memunculkan pesan peringatan field wajib diisi.", "status": "Valid"},
    {"id": "TC-08", "modul": "Dropdown Cascading", "skenario": "Pemilihan nama unit sekolah pada form master lokasi", "hasil_diharapkan": "Dropdown ruangan otomatis memuat daftar ruangan yang hanya berada di unit tersebut.", "status": "Valid"},
    {"id": "TC-09", "modul": "Upload Bukti Foto", "skenario": "Unggah foto kondisi fasilitas rusak berukuran 5 MB via smartphone", "hasil_diharapkan": "Sistem mengompresi gambar otomatis via Canvas menjadi < 400 KB sebelum dikirim ke API.", "status": "Valid"},
    {"id": "TC-10", "modul": "Akses Kamera", "skenario": "Petugas lapangan menekan tombol jepret foto pada ponsel", "hasil_diharapkan": "Peramban mobile otomatis membuka kamera perangkat belakang (environment capture).", "status": "Valid"},
    {"id": "TC-11", "modul": "Review ASP", "skenario": "ASP Wilayah memverifikasi fisik dan menyetujui tiket pengajuan", "hasil_diharapkan": "Status tiket berubah menjadi 'ASP Reviewed', estimasi biaya tersimpan, tercatat di audit log.", "status": "Valid"},
    {"id": "TC-12", "modul": "Review ASP", "skenario": "ASP Wilayah meminta perbaikan form (Revision Required) ke Kepala Sekolah", "hasil_diharapkan": "Status tiket menjadi 'Revision Required', catatan revisi muncul pada akun Kepala Sekolah.", "status": "Valid"},
    {"id": "TC-13", "modul": "Approval KD Sarpras", "skenario": "KD Sarpras menyetujui tiket yang telah diverifikasi ASP", "hasil_diharapkan": "Status tiket berubah menjadi 'KD Approved', alokasi anggaran tercatat, tiket siap diteruskan ke Admin PO.", "status": "Valid"},
    {"id": "TC-14", "modul": "Approval KD Sarpras", "skenario": "KD Sarpras menolak tiket pengajuan dengan alasan tidak mendesak", "hasil_diharapkan": "Status tiket berubah menjadi 'Rejected', alasan penolakan tercatat di linimasa audit log.", "status": "Valid"},
    {"id": "TC-15", "modul": "Purchase Order", "skenario": "Admin Sarpras membuat Purchase Order (PO) dari tiket approved", "hasil_diharapkan": "Nomor PO otomatis terbit, terhubung ke vendor rekanan, total nominal terhitung akurat.", "status": "Valid"},
    {"id": "TC-16", "modul": "Terbilang Rupiah", "skenario": "Sistem mencetak nominal PO senilai Rp 12.500.000,- ke dokumen", "hasil_diharapkan": "Helper otomatis mengonversi menjadi 'Dua Belas Juta Lima Ratus Ribu Rupiah'.", "status": "Valid"},
    {"id": "TC-17", "modul": "Cetak Dokumen PDF", "skenario": "Admin/Kepsek mengklik cetak Berita Acara Serah Terima (BAST)", "hasil_diharapkan": "Sistem menghasilkan file PDF BAST berkop resmi YGPKB dan tabel rincian rapi.", "status": "Valid"},
    {"id": "TC-18", "modul": "Timeline Tracker", "skenario": "Pengguna membuka halaman detail tiket pengajuan", "hasil_diharapkan": "Komponen linimasa visual menampilkan urutan progres status, tanggal, dan nama aktor secara runtut.", "status": "Valid"},
    {"id": "TC-19", "modul": "Responsivitas UI", "skenario": "Pengguna mengakses aplikasi dari layar smartphone (< 768px)", "hasil_diharapkan": "Sidebar berubah menjadi drawer tersembunyi yang dapat dibuka via tombol hamburger menu.", "status": "Valid"},
    {"id": "TC-20", "modul": "Export Laporan", "skenario": "Admin Sarpras mengekspor rekapitulasi tiket bulanan ke Excel", "hasil_diharapkan": "Sistem mengunduh berkas laporan tiket dalam format .xlsx dengan data terfilter akurat.", "status": "Valid"}
]


def create_excel_logbook(filepath):
    print(f"Creating Excel logbook at {filepath}...")
    wb = openpyxl.Workbook()
    
    # Styles
    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Calibri", size=11, bold=True, color="475569")
    font_section = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=9, color="1E293B")
    font_bold_data = Font(name="Calibri", size=9, bold=True, color="1E293B")
    
    fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_blue_light = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_success = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_success = Font(name="Calibri", size=9, bold=True, color="166534")

    border_thin = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )
    border_header = Border(
        left=Side(style='thin', color="1E3A8A"),
        right=Side(style='thin', color="1E3A8A"),
        top=Side(style='thin', color="1E3A8A"),
        bottom=Side(style='medium', color="0F172A")
    )

    # ==================== SHEET 1: PROFIL PROYEK ====================
    ws_info = wb.active
    ws_info.title = "Profil Proyek"
    ws_info.views.sheetView[0].showGridLines = True
    
    ws_info.merge_cells("A1:G1")
    ws_info["A1"] = "LOGBOOK PENGEMBANGAN SISTEM INFORMASI SARANA & PRASARANA (SIPARAS)"
    ws_info["A1"].font = font_title
    ws_info["A1"].alignment = Alignment(vertical="center")

    ws_info.merge_cells("A2:G2")
    ws_info["A2"] = "Yayasan Gereja Protestan Kampung Bali (YGPKB) — Pontianak & Kubu Raya, Kalimantan Barat"
    ws_info["A2"].font = font_subtitle
    ws_info["A2"].alignment = Alignment(vertical="center")
    
    ws_info.row_dimensions[1].height = 25
    ws_info.row_dimensions[2].height = 20
    ws_info.row_dimensions[4].height = 20

    profile_data = [
        ("Nama Sistem", "SIPARAS (Sistem Informasi Sarana dan Prasarana)"),
        ("Instansi / Lembaga", "Yayasan Gereja Protestan Kampung Bali (YGPKB)"),
        ("Cakupan Operasional", "14 Unit Kerja (12 Sekolah TK-SMA, Kampus ITBSS, dan Kantor Pusat Yayasan)"),
        ("Pembagian Wilayah ASP", "5 Wilayah Kerja (Kubu Raya, Reguler, IBC, Kota, ITBSS)"),
        ("Arsitektur Sistem", "Monorepo Web Application (Backend API & Frontend Client)"),
        ("Teknologi Backend", "PHP 8.3/8.4, Laravel 11, Laravel Sanctum, DomPDF, Laravel Excel, MySQL 8.4 LTS"),
        ("Teknologi Frontend", "Vue.js 3 (Composition API), Vite, Pinia, Tailwind CSS, Axios, HTML5 Canvas Compressor"),
        ("Model Pengembangan", "Software Development Life Cycle (SDLC) - Waterfall / Iterative Agile Model"),
        ("Periode Pengerjaan", "10 Agustus 2026 s/d 22 September 2026 (7 Minggu / 31 Hari Kerja)"),
        ("Peran Pengguna (7 Roles)", "1. Kepala Sekolah/Pimpinan ITBSS, 2. ASP Wilayah, 3. KD Sarpras, 4. Admin Sarpras, 5. Koordinator Utama/SKI, 6. Dept Kantor Pusat, 7. Super Admin"),
        ("5 Prosedur Utama", "1. Perbaikan (Corrective), 2. Pemeliharaan (Preventive), 3. Pengadaan (Procurement), 4. Mutasi Barang (Asset Transfer), 5. Pemutihan Aset (Write-off)"),
        ("Tautan Demo Live", "https://siparas-frontend.vercel.app")
    ]

    r = 4
    ws_info.cell(row=r, column=1, value="INFORMASI UMUM PROYEK").font = font_section
    r += 1
    
    for item, val in profile_data:
        ws_info.row_dimensions[r].height = 22
        c1 = ws_info.cell(row=r, column=1, value=item)
        c1.font = font_bold_data
        c1.fill = fill_blue_light
        c1.border = border_thin
        c1.alignment = Alignment(vertical="center")

        ws_info.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c2 = ws_info.cell(row=r, column=2, value=val)
        c2.font = font_data
        c2.border = border_thin
        c2.alignment = Alignment(vertical="center")
        
        # Bordering merged cells
        for col_idx in range(3, 8):
            ws_info.cell(row=r, column=col_idx).border = border_thin
        r += 1

    ws_info.column_dimensions["A"].width = 25
    for col_let in ["B", "C", "D", "E", "F", "G"]:
        ws_info.column_dimensions[col_let].width = 18

    # ==================== SHEET 2: LOGBOOK HARIAN ====================
    ws_daily = wb.create_sheet(title="Logbook Harian")
    ws_daily.views.sheetView[0].showGridLines = True
    
    ws_daily.merge_cells("A1:H1")
    ws_daily["A1"] = "LOGBOOK HARIAN PENGERJAAN SISTEM SIPARAS"
    ws_daily["A1"].font = font_title
    ws_daily.row_dimensions[1].height = 25

    ws_daily.merge_cells("A2:H2")
    ws_daily["A2"] = "Catatan Aktivitas Rekayasa Perangkat Lunak (10 Agustus 2026 - 22 September 2026)"
    ws_daily["A2"].font = font_subtitle
    ws_daily.row_dimensions[2].height = 18
    ws_daily.row_dimensions[3].height = 10

    headers_daily = ["No", "Hari / Tanggal", "Jam Kerja", "Tahapan SDLC", "Modul / Fitur", "Rincian Aktivitas Pengerjaan", "Output / Hasil", "Status"]
    header_row = 4
    ws_daily.row_dimensions[header_row].height = 28
    
    for c_idx, h in enumerate(headers_daily, 1):
        cell = ws_daily.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_header

    for idx, log in enumerate(DAILY_LOGS, start=5):
        ws_daily.row_dimensions[idx].height = 55
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra if is_even else None
        
        c1 = ws_daily.cell(row=idx, column=1, value=log["no"])
        c1.alignment = Alignment(horizontal="center", vertical="center")
        
        c2 = ws_daily.cell(row=idx, column=2, value=log["hari_tanggal"])
        c2.alignment = Alignment(horizontal="center", vertical="center")
        
        c3 = ws_daily.cell(row=idx, column=3, value=log["jam"])
        c3.alignment = Alignment(horizontal="center", vertical="center")
        
        c4 = ws_daily.cell(row=idx, column=4, value=log["tahap"])
        c4.alignment = Alignment(vertical="center", wrap_text=True)
        
        c5 = ws_daily.cell(row=idx, column=5, value=log["modul"])
        c5.alignment = Alignment(vertical="center", wrap_text=True)
        c5.font = font_bold_data
        
        c6 = ws_daily.cell(row=idx, column=6, value=log["kegiatan"])
        c6.alignment = Alignment(vertical="center", wrap_text=True)
        
        c7 = ws_daily.cell(row=idx, column=7, value=log["output"])
        c7.alignment = Alignment(vertical="center", wrap_text=True)
        
        c8 = ws_daily.cell(row=idx, column=8, value=log["status"])
        c8.alignment = Alignment(horizontal="center", vertical="center")
        c8.fill = fill_success
        c8.font = font_success

        for c_idx in range(1, 9):
            cell = ws_daily.cell(row=idx, column=c_idx)
            cell.border = border_thin
            if c_idx != 8 and row_fill:
                cell.fill = row_fill
            if c_idx not in [5, 8]:
                cell.font = font_data

    daily_widths = [6, 18, 15, 24, 25, 48, 35, 12]
    for c_idx, w in enumerate(daily_widths, 1):
        col_letter = get_column_letter(c_idx)
        ws_daily.column_dimensions[col_letter].width = w

    # ==================== SHEET 3: REKAPITULASI MINGGUAN ====================
    ws_weekly = wb.create_sheet(title="Rekapitulasi Mingguan")
    ws_weekly.views.sheetView[0].showGridLines = True
    
    ws_weekly.merge_cells("A1:G1")
    ws_weekly["A1"] = "REKAPITULASI PROGRES MINGGUAN & EVALUASI PEMBIMBING"
    ws_weekly["A1"].font = font_title
    ws_weekly.row_dimensions[1].height = 25

    ws_weekly.merge_cells("A2:G2")
    ws_weekly["A2"] = "Evaluasi Capaian Milestone, Kendala Teknis, dan Solusi Implementasi"
    ws_weekly["A2"].font = font_subtitle
    ws_weekly.row_dimensions[2].height = 18
    ws_weekly.row_dimensions[3].height = 10

    headers_weekly = ["Minggu", "Rentang Tanggal", "Fokus Tahapan", "Capaian & Milestone", "Kendala Teknis Dihadapi", "Solusi & Tindak Lanjut", "Paraf/Status"]
    header_row = 4
    ws_weekly.row_dimensions[header_row].height = 28
    
    for c_idx, h in enumerate(headers_weekly, 1):
        cell = ws_weekly.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_header

    for idx, wlog in enumerate(WEEKLY_LOGS, start=5):
        ws_weekly.row_dimensions[idx].height = 65
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra if is_even else None
        
        c1 = ws_weekly.cell(row=idx, column=1, value=wlog["minggu"])
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = font_bold_data
        
        c2 = ws_weekly.cell(row=idx, column=2, value=wlog["tanggal"])
        c2.alignment = Alignment(horizontal="center", vertical="center")
        
        c3 = ws_weekly.cell(row=idx, column=3, value=wlog["fokus"])
        c3.alignment = Alignment(vertical="center", wrap_text=True)
        c3.font = font_bold_data
        
        c4 = ws_weekly.cell(row=idx, column=4, value=wlog["capaian"])
        c4.alignment = Alignment(vertical="center", wrap_text=True)
        
        c5 = ws_weekly.cell(row=idx, column=5, value=wlog["kendala"])
        c5.alignment = Alignment(vertical="center", wrap_text=True)
        
        c6 = ws_weekly.cell(row=idx, column=6, value=wlog["solusi"])
        c6.alignment = Alignment(vertical="center", wrap_text=True)
        
        c7 = ws_weekly.cell(row=idx, column=7, value=wlog["paraf"])
        c7.alignment = Alignment(horizontal="center", vertical="center")
        c7.fill = fill_success
        c7.font = font_success

        for c_idx in range(1, 8):
            cell = ws_weekly.cell(row=idx, column=c_idx)
            cell.border = border_thin
            if c_idx != 7 and row_fill:
                cell.fill = row_fill
            if c_idx not in [1, 3, 7]:
                cell.font = font_data

    weekly_widths = [12, 18, 25, 38, 32, 35, 14]
    for c_idx, w in enumerate(weekly_widths, 1):
        col_letter = get_column_letter(c_idx)
        ws_weekly.column_dimensions[col_letter].width = w

    # ==================== SHEET 4: BLACKBOX TESTING ====================
    ws_testing = wb.create_sheet(title="Matriks Testing (QA)")
    ws_testing.views.sheetView[0].showGridLines = True
    
    ws_testing.merge_cells("A1:E1")
    ws_testing["A1"] = "MATRIKS PENGUJIAN BLACKBOX TESTING SISTEM SIPARAS"
    ws_testing["A1"].font = font_title
    ws_testing.row_dimensions[1].height = 25

    ws_testing.merge_cells("A2:E2")
    ws_testing["A2"] = "Hasil Pengujian Fungsionalitas 7 Role Pengguna, RBAC, dan Fitur Esensial"
    ws_testing["A2"].font = font_subtitle
    ws_testing.row_dimensions[2].height = 18
    ws_testing.row_dimensions[3].height = 10

    headers_test = ["ID Uji", "Modul / Fitur", "Skenario Pengujian", "Hasil yang Diharapkan", "Status"]
    header_row = 4
    ws_testing.row_dimensions[header_row].height = 28
    
    for c_idx, h in enumerate(headers_test, 1):
        cell = ws_testing.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_header

    for idx, tc in enumerate(TEST_CASES, start=5):
        ws_testing.row_dimensions[idx].height = 35
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra if is_even else None
        
        c1 = ws_testing.cell(row=idx, column=1, value=tc["id"])
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = font_bold_data
        
        c2 = ws_testing.cell(row=idx, column=2, value=tc["modul"])
        c2.alignment = Alignment(vertical="center")
        c2.font = font_bold_data
        
        c3 = ws_testing.cell(row=idx, column=3, value=tc["skenario"])
        c3.alignment = Alignment(vertical="center", wrap_text=True)
        
        c4 = ws_testing.cell(row=idx, column=4, value=tc["hasil_diharapkan"])
        c4.alignment = Alignment(vertical="center", wrap_text=True)
        
        c5 = ws_testing.cell(row=idx, column=5, value=tc["status"])
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c5.fill = fill_success
        c5.font = font_success

        for c_idx in range(1, 6):
            cell = ws_testing.cell(row=idx, column=c_idx)
            cell.border = border_thin
            if c_idx != 5 and row_fill:
                cell.fill = row_fill
            if c_idx not in [1, 2, 5]:
                cell.font = font_data

    test_widths = [10, 22, 42, 48, 12]
    for c_idx, w in enumerate(test_widths, 1):
        col_letter = get_column_letter(c_idx)
        ws_testing.column_dimensions[col_letter].width = w

    wb.save(filepath)
    print(f"Successfully generated {filepath}")


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins for a docx cell in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    """Set cell background shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def create_word_logbook(filepath):
    print(f"Creating Word logbook at {filepath}...")
    doc = docx.Document()
    
    # Set standard margins (1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.different_first_page_header_footer = False

    # Title & Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("LOGBOOK PENGERJAAN KERJA PRAKTEK / PROYEK\nSISTEM INFORMASI SARANA DAN PRASARANA (SIPARAS)")
    run_title.bold = True
    run_title.font.size = Pt(14)
    run_title.font.name = "Arial"
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Yayasan Gereja Protestan Kampung Bali (YGPKB) — Pontianak & Kubu Raya\nPeriode: 10 Agustus 2026 s/d 22 September 2026")
    run_sub.font.size = Pt(10.5)
    run_sub.font.name = "Arial"
    run_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph() # Spacing

    # Section 1: Informasi Proyek
    h1 = doc.add_heading("I. IDENTITAS PROYEK & SPESIFIKASI SISTEM", level=2)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(4)
    
    tbl_info = doc.add_table(rows=6, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_rows = [
        ("Nama Sistem", "SIPARAS (Sistem Informasi Sarana & Prasarana)"),
        ("Instansi / Organisasi Mitra", "Yayasan Gereja Protestan Kampung Bali (YGPKB)"),
        ("Cakupan 14 Unit Kerja", "12 Sekolah (TK, SD, SMP, SMK, SMA Reguler & IBC), Kampus ITBSS, dan Kantor Pusat Yayasan"),
        ("Arsitektur & Tech Stack", "Monorepo: Laravel 11 Backend (PHP 8.3, Sanctum, DomPDF, Excel, MySQL 8.4) & Vue 3 Frontend (Vite, Pinia, Tailwind CSS)"),
        ("7 Hak Akses Peran (POV)", "Kepala Sekolah (13 akun unit), ASP Wilayah (5 wilayah), KD Sarpras, Admin Sarpras, Koordinator, Dept YGPKB, Super Admin"),
        ("Model Pengembangan", "SDLC - Waterfall / Iterative Agile (Analisis -> Desain -> Implementasi -> Pengujian -> Deployment)")
    ]

    for idx, (label, val) in enumerate(info_rows):
        row = tbl_info.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.name = "Arial"
        set_cell_shading(c0, "F1F5F9")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.name = "Arial"
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)

    doc.add_paragraph() # Spacing

    # Section 2: Logbook Harian
    h2 = doc.add_heading("II. CATATAN AKTIVITAS HARIAN (DAILY LOGBOOK)", level=2)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)

    p_intro = doc.add_paragraph()
    r_intro = p_intro.add_run("Tabel berikut merekam aktivitas teknis harian pengembangan sistem SIPARAS selama 31 hari kerja efektif:")
    r_intro.font.size = Pt(9.5)
    r_intro.font.italic = True
    r_intro.font.name = "Arial"

    tbl_daily = doc.add_table(rows=len(DAILY_LOGS) + 1, cols=6)
    tbl_daily.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers_doc = ["No", "Hari, Tanggal & Jam", "Modul / Fitur", "Rincian Aktivitas Teknis", "Output / Hasil", "Paraf"]
    widths_doc = [Inches(0.4), Inches(1.3), Inches(1.3), Inches(2.3), Inches(1.6), Inches(0.6)]

    # Header Row
    hdr_row = tbl_daily.rows[0]
    for c_idx, text in enumerate(headers_doc):
        cell = hdr_row.cells[c_idx]
        cell.width = widths_doc[c_idx]
        set_cell_shading(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data Rows
    for idx, log in enumerate(DAILY_LOGS, start=1):
        row = tbl_daily.rows[idx]
        is_even = (idx % 2 == 0)
        shd_color = "F8FAFC" if is_even else "FFFFFF"

        # Col 0: No
        c0 = row.cells[0]
        c0.width = widths_doc[0]
        set_cell_shading(c0, shd_color)
        set_cell_margins(c0, top=60, bottom=60, left=60, right=60)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p0.add_run(str(log["no"]))
        r.font.size = Pt(8.5)
        r.font.name = "Arial"

        # Col 1: Tanggal & Jam
        c1 = row.cells[1]
        c1.width = widths_doc[1]
        set_cell_shading(c1, shd_color)
        set_cell_margins(c1, top=60, bottom=60, left=60, right=60)
        p1 = c1.paragraphs[0]
        r1a = p1.add_run(log["hari_tanggal"] + "\n")
        r1a.bold = True
        r1a.font.size = Pt(8.5)
        r1a.font.name = "Arial"
        r1b = p1.add_run(log["jam"])
        r1b.font.size = Pt(8)
        r1b.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        r1b.font.name = "Arial"

        # Col 2: Modul & Tahap
        c2 = row.cells[2]
        c2.width = widths_doc[2]
        set_cell_shading(c2, shd_color)
        set_cell_margins(c2, top=60, bottom=60, left=60, right=60)
        p2 = c2.paragraphs[0]
        r2a = p2.add_run(log["modul"] + "\n")
        r2a.bold = True
        r2a.font.size = Pt(8.5)
        r2a.font.name = "Arial"
        r2b = p2.add_run(log["tahap"])
        r2b.font.size = Pt(7.5)
        r2b.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        r2b.font.name = "Arial"

        # Col 3: Aktivitas Teknis
        c3 = row.cells[3]
        c3.width = widths_doc[3]
        set_cell_shading(c3, shd_color)
        set_cell_margins(c3, top=60, bottom=60, left=60, right=60)
        p3 = c3.paragraphs[0]
        r3 = p3.add_run(log["kegiatan"])
        r3.font.size = Pt(8.5)
        r3.font.name = "Arial"

        # Col 4: Output / Hasil
        c4 = row.cells[4]
        c4.width = widths_doc[4]
        set_cell_shading(c4, shd_color)
        set_cell_margins(c4, top=60, bottom=60, left=60, right=60)
        p4 = c4.paragraphs[0]
        r4 = p4.add_run(log["output"])
        r4.font.size = Pt(8.5)
        r4.font.name = "Arial"

        # Col 5: Paraf / Status
        c5 = row.cells[5]
        c5.width = widths_doc[5]
        set_cell_shading(c5, shd_color)
        set_cell_margins(c5, top=60, bottom=60, left=60, right=60)
        p5 = c5.paragraphs[0]
        p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r5 = p5.add_run("OK")
        r5.bold = True
        r5.font.size = Pt(8.5)
        r5.font.name = "Arial"
        r5.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

    doc.add_page_break()

    # Section 3: Rekapitulasi Mingguan
    h3 = doc.add_heading("III. REKAPITULASI MINGGUAN & EVALUASI", level=2)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)

    tbl_week = doc.add_table(rows=len(WEEKLY_LOGS) + 1, cols=5)
    tbl_week.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_w_doc = ["Minggu / Tanggal", "Fokus Tahapan", "Capaian & Hasil", "Kendala Teknis & Solusi", "Paraf"]
    widths_w_doc = [Inches(1.3), Inches(1.4), Inches(2.2), Inches(2.2), Inches(0.6)]

    # Header
    for c_idx, text in enumerate(headers_w_doc):
        cell = tbl_week.rows[0].cells[c_idx]
        cell.width = widths_w_doc[c_idx]
        set_cell_shading(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for idx, wlog in enumerate(WEEKLY_LOGS, start=1):
        row = tbl_week.rows[idx]
        is_even = (idx % 2 == 0)
        shd_color = "F8FAFC" if is_even else "FFFFFF"

        # Col 0: Minggu
        c0 = row.cells[0]
        c0.width = widths_w_doc[0]
        set_cell_shading(c0, shd_color)
        set_cell_margins(c0, top=60, bottom=60, left=60, right=60)
        p0 = c0.paragraphs[0]
        r0a = p0.add_run(wlog["minggu"] + "\n")
        r0a.bold = True
        r0a.font.size = Pt(8.5)
        r0a.font.name = "Arial"
        r0b = p0.add_run(wlog["tanggal"])
        r0b.font.size = Pt(8)
        r0b.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        r0b.font.name = "Arial"

        # Col 1: Fokus
        c1 = row.cells[1]
        c1.width = widths_w_doc[1]
        set_cell_shading(c1, shd_color)
        set_cell_margins(c1, top=60, bottom=60, left=60, right=60)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(wlog["fokus"])
        r1.bold = True
        r1.font.size = Pt(8.5)
        r1.font.name = "Arial"

        # Col 2: Capaian
        c2 = row.cells[2]
        c2.width = widths_w_doc[2]
        set_cell_shading(c2, shd_color)
        set_cell_margins(c2, top=60, bottom=60, left=60, right=60)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(wlog["capaian"])
        r2.font.size = Pt(8.5)
        r2.font.name = "Arial"

        # Col 3: Kendala & Solusi
        c3 = row.cells[3]
        c3.width = widths_w_doc[3]
        set_cell_shading(c3, shd_color)
        set_cell_margins(c3, top=60, bottom=60, left=60, right=60)
        p3 = c3.paragraphs[0]
        r3a = p3.add_run("Kendala: ")
        r3a.bold = True
        r3a.font.size = Pt(8)
        r3a.font.name = "Arial"
        r3b = p3.add_run(wlog["kendala"] + "\n\n")
        r3b.font.size = Pt(8)
        r3b.font.name = "Arial"
        
        r3c = p3.add_run("Solusi: ")
        r3c.bold = True
        r3c.font.size = Pt(8)
        r3c.font.name = "Arial"
        r3d = p3.add_run(wlog["solusi"])
        r3d.font.size = Pt(8)
        r3d.font.name = "Arial"

        # Col 4: Paraf
        c4 = row.cells[4]
        c4.width = widths_w_doc[4]
        set_cell_shading(c4, shd_color)
        set_cell_margins(c4, top=60, bottom=60, left=60, right=60)
        p4 = c4.paragraphs[0]
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r4 = p4.add_run("Valid")
        r4.bold = True
        r4.font.size = Pt(8.5)
        r4.font.name = "Arial"
        r4.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

    doc.add_paragraph() # Spacing

    # Section 4: Lembar Pengesahan
    h4 = doc.add_heading("IV. LEMBAR PENGESAHAN LOGBOOK KERJA PRAKTEK", level=2)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(8)

    p_valid = doc.add_paragraph()
    r_valid = p_valid.add_run("Demikian logbook pengerjaan ini disusun secara benar berdasarkan pelaksanaan kegiatan perancangan, implementasi, dan pengujian Sistem Informasi Sarana dan Prasarana (SIPARAS) pada Yayasan Gereja Protestan Kampung Bali (YGPKB).")
    r_valid.font.size = Pt(9.5)
    r_valid.font.name = "Arial"

    tbl_ttd = doc.add_table(rows=1, cols=3)
    tbl_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    col_w = Inches(2.4)
    for c_idx in range(3):
        tbl_ttd.rows[0].cells[c_idx].width = col_w

    cell_mhs = tbl_ttd.rows[0].cells[0]
    p_m = cell_mhs.paragraphs[0]
    p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_m.add_run("Pontianak, 22 September 2026\nMahasiswa Pelaksana,\n\n\n\n\n( ............................................ )\nNIM : ...................................").font.size = Pt(9)

    cell_mitra = tbl_ttd.rows[0].cells[1]
    p_mitra = cell_mitra.paragraphs[0]
    p_mitra.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mitra.add_run("Mengetahui,\nPembimbing Lapangan (YGPKB),\n\n\n\n\n( ............................................ )\nKD Sarpras YGPKB").font.size = Pt(9)

    cell_dosen = tbl_ttd.rows[0].cells[2]
    p_dos = cell_dosen.paragraphs[0]
    p_dos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dos.add_run("Menyetujui,\nDosen Pembimbing Kerja Praktek,\n\n\n\n\n( ............................................ )\nNIDN : ...................................").font.size = Pt(9)

    doc.save(filepath)
    print(f"Successfully generated {filepath}")


if __name__ == "__main__":
    excel_path = os.path.abspath("Logbook_Pengerjaan_SIPARAS.xlsx")
    docx_path = os.path.abspath("Logbook_Pengerjaan_SIPARAS.docx")
    create_excel_logbook(excel_path)
    create_word_logbook(docx_path)
    print("Done! Both files created successfully.")
